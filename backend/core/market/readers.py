"""Read market tables as they are: CSV (encoding and delimiter detected) and XLSX (every sheet,
merged ranges filled from their first cell, dates as ISO text). Nothing is interpreted here: the
cells land in ``market_imports.raw`` untouched and :mod:`core.market.rules` /
:mod:`core.market.llm_map` say what they mean."""

from __future__ import annotations

import csv
import hashlib
import io
from datetime import date, datetime, time
from pathlib import PurePath
from typing import Any

from core.market.model import Cell, RawTable, Sheet

CSV_EXTENSIONS = (".csv", ".tsv", ".txt")
XLSX_EXTENSIONS = (".xlsx", ".xlsm")
DELIMITERS = (";", ",", "\t", "|")
ENCODINGS = ("utf-8-sig", "cp1250")  # then latin-1, which decodes anything


class ReadError(ValueError):
    """The file cannot be read as a table (the message says why, for the uploader)."""


def sha256_hex(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _cell(value: Any) -> Cell:
    if value is None:
        return None
    if isinstance(value, bool):
        return str(value).lower()
    if isinstance(value, int | float):
        return value
    if isinstance(value, datetime):
        return (value.date() if value.time() == time() else value).isoformat()
    if isinstance(value, date | time):
        return value.isoformat()
    text = str(value).strip()
    return text or None


def _trim(rows: list[list[Cell]]) -> list[list[Cell]]:
    """Drop trailing empty cells of each row and trailing empty rows."""
    out = []
    for row in rows:
        end = len(row)
        while end and row[end - 1] is None:
            end -= 1
        out.append(row[:end])
    while out and not out[-1]:
        out.pop()
    return out


def decode_text(data: bytes) -> tuple[str, str]:
    for encoding in ENCODINGS:
        try:
            return data.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1"), "latin-1"


def sniff_delimiter(text: str) -> str:
    """The delimiter that splits most sample lines into the same number of fields."""
    lines = [line for line in text.splitlines() if line.strip()][:30]
    best, best_score = DELIMITERS[0], (-1, -1)
    for delimiter in DELIMITERS:
        counts = [len(r) for r in csv.reader(lines, delimiter=delimiter)]
        split = [c for c in counts if c > 1]
        if not split:
            continue
        mode = max(set(split), key=split.count)
        score = (sum(1 for c in split if c == mode), mode)
        if score > best_score:
            best, best_score = delimiter, score
    return best


def read_csv(data: bytes, *, max_rows: int = 5000, name: str = "csv") -> RawTable:
    text, encoding = decode_text(data)
    delimiter = sniff_delimiter(text)
    rows = [[_cell(c) for c in row] for row in csv.reader(io.StringIO(text), delimiter=delimiter)]
    rows = _trim(rows)
    if not rows:
        raise ReadError("the file holds no rows")
    if len(rows) > max_rows:
        raise ReadError(f"{len(rows)} rows: more than the {max_rows} a market table may have")
    return RawTable(
        format="csv", encoding=encoding, delimiter=delimiter, sheets=[Sheet(name, rows)]
    )


def read_xlsx(data: bytes, *, max_rows: int = 5000) -> RawTable:
    try:
        from openpyxl import load_workbook
        from openpyxl.utils import range_boundaries
    except ImportError as exc:  # pragma: no cover - a core dependency
        raise ReadError("openpyxl is not installed") from exc
    try:
        workbook = load_workbook(io.BytesIO(data), data_only=True)
    except Exception as exc:  # noqa: BLE001 - a corrupt or foreign file: say so
        raise ReadError(f"not a readable .xlsx workbook ({type(exc).__name__})") from exc
    sheets: list[Sheet] = []
    for ws in workbook.worksheets:
        if ws.max_row > max_rows:
            raise ReadError(
                f"sheet {ws.title!r} has {ws.max_row} rows: more than the {max_rows} allowed"
            )
        rows = [[_cell(c) for c in row] for row in ws.iter_rows(values_only=True)]
        merged = []
        for merged_range in ws.merged_cells.ranges:
            min_col, min_row, max_col, max_row = range_boundaries(str(merged_range))
            merged.append(str(merged_range))
            first = rows[min_row - 1][min_col - 1] if min_row - 1 < len(rows) else None
            for r in range(min_row - 1, min(max_row, len(rows))):
                row = rows[r]
                while len(row) < max_col:
                    row.append(None)
                for c in range(min_col - 1, max_col):
                    row[c] = first
        rows = _trim(rows)
        if rows:
            sheets.append(Sheet(ws.title, rows, merged=sorted(merged)))
    if not sheets:
        raise ReadError("the workbook holds no data")
    return RawTable(format="xlsx", sheets=sheets)


def read_table(filename: str, data: bytes, *, max_rows: int = 5000) -> RawTable:
    """Read an uploaded market table by its extension (and signature)."""
    suffix = PurePath(filename).suffix.lower()
    if suffix in XLSX_EXTENSIONS or data[:4] == b"PK\x03\x04":
        return read_xlsx(data, max_rows=max_rows)
    if suffix in CSV_EXTENSIONS:
        return read_csv(data, max_rows=max_rows, name=PurePath(filename).stem or "csv")
    if suffix == ".xls":
        raise ReadError("old .xls workbooks are not read: save the sheet as .xlsx or .csv")
    raise ReadError(f"unsupported file type {suffix or '(none)'}: upload .csv or .xlsx")
